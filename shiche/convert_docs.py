#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
将 shiche 文件夹中的 docx 文档转换为 Markdown。
用法: python3 convert_docs.py
"""
import docx
import re


def normalize_style(style_name: str) -> str:
    """归一化样式名"""
    if style_name is None:
        return "Normal"
    return style_name.strip()


def heading_level(style_name: str):
    m = re.match(r"Heading\s*(\d)", style_name or "")
    if m:
        return int(m.group(1))
    return None


def convert_docx(path: str) -> str:
    d = docx.Document(path)
    lines = []
    in_list = False

    def close_list():
        nonlocal in_list
        if in_list:
            lines.append("")
            in_list = False

    # 先处理段落
    for p in d.paragraphs:
        text = p.text.strip()
        style = normalize_style(p.style.name)

        if not text:
            close_list()
            continue

        # 代码块（等宽样式 / 含命令的样式）
        if style in ("Plain Text", "代码", "Code") or (style == "Normal" and text.startswith(("ssh ", "cd ", "aem ", "sudo ", "whoami", "hostname", "pwd", "Get-", "Add-", "ssh -V"))):
            close_list()
            lines.append(f"```\n{text}\n```")
            continue

        lvl = heading_level(style)
        if lvl:
            close_list()
            lines.append(f"{'#' * min(lvl + 1, 6)} {text}")
            continue

        if style == "List Bullet" or text.startswith(("•", "- ")) and style in ("List Bullet",):
            if not in_list:
                lines.append("")
                in_list = True
            lines.append(f"- {text.lstrip('•- ')}")
            continue

        if text.startswith(">"):
            close_list()
            lines.append(text)
            continue

        close_list()
        lines.append(text)

    close_list()

    # 处理表格
    for ti, t in enumerate(d.tables):
        rows = []
        for row in t.rows:
            rows.append([c.text.strip().replace("\n", " ") for c in row.cells])
        if not rows:
            continue
        # 生成 markdown 表格
        lines.append("")
        header = rows[0]
        lines.append("| " + " | ".join(header) + " |")
        lines.append("| " + " | ".join(["---"] * len(header)) + " |")
        for r in rows[1:]:
            lines.append("| " + " | ".join(r) + " |")
        lines.append("")

    return "\n".join(lines)


def main():
    jobs = [
        ("附录一：智能驾驶实车赛--笔记本远程实车连接操作手册-0(1).docx",
         "附录一_智能驾驶实车赛_笔记本远程实车连接操作手册.md"),
        ("自动驾驶闭环文档.docx", "自动驾驶闭环文档.md"),
    ]
    for src, dst in jobs:
        md = convert_docx(src)
        with open(dst, "w", encoding="utf-8") as f:
            f.write(md)
        print(f"✅ {src} -> {dst} ({len(md)} 字符)")


if __name__ == "__main__":
    main()
